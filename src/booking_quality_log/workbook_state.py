"""Read back a prior "Booking Quality Log" workbook's Reservation IDs and
hand-typed manual columns, so a rerun can:

1. tell whether any brand-new confirmed booking has shown up since that
   run (the daily skip-if-nothing-new check), and
2. carry every still-visible booking's manual notes forward untouched.

Both needs read the same data, so one function serves both.

Every column is located by reading the prior file's OWN header row and
matching header text, not a hardcoded or relative column number: this is
deliberate. "Reservation ID" never changes, so it anchors the lookup even
across a layout change. Each of the *current* tool's manual column names
(workbook_build.MANUAL_COLUMNS) is then looked up by that same text in the
OLD file -- a column present in the old file but no longer manual (e.g.
"LY Weekday Occ." once it became the computed "LY occ." column) is
correctly left behind rather than misread into the wrong slot, and a
manual column that's brand new in this version just comes back blank for
an older file that never had it.
"""

from __future__ import annotations

from pathlib import Path

import openpyxl

from .workbook_build import HEADER_ROW, MANUAL_COLUMNS, SHEET_NAME

RESERVATION_ID_HEADER = "Reservation ID"


def load_prior_reservation_state(path: Path) -> dict[str, tuple]:
    """Returns {reservation_id: (manual-column values, in MANUAL_COLUMNS
    order)} for every row in the prior file's "Booking Quality Log" sheet
    -- including rows where every manual value is still blank, since the
    keys alone are what the daily skip check needs.
    """
    if not path.exists():
        return {}
    wb = openpyxl.load_workbook(path, data_only=False)
    if SHEET_NAME not in wb.sheetnames:
        return {}
    ws = wb[SHEET_NAME]

    header_row = ws[HEADER_ROW]
    header_index: dict[str, int] = {
        cell.value: i for i, cell in enumerate(header_row) if cell.value
    }
    reservation_id_col = header_index.get(RESERVATION_ID_HEADER)
    if reservation_id_col is None:
        return {}
    manual_cols = [header_index.get(name) for name in MANUAL_COLUMNS]

    result: dict[str, tuple] = {}
    for row in ws.iter_rows(min_row=HEADER_ROW + 1):
        if len(row) <= reservation_id_col:
            continue
        reservation_id = row[reservation_id_col].value
        if not reservation_id:
            continue
        manual_values = tuple(
            row[col].value if col is not None and len(row) > col else None for col in manual_cols
        )
        result[str(reservation_id)] = manual_values
    return result
