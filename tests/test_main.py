import datetime as dt
from unittest.mock import MagicMock, patch

import openpyxl

from booking_quality_log import main as bql_main
from booking_quality_log.config import Listing
from booking_quality_log.workbook_build import MANUAL_COLS_START, MANUAL_COLUMNS

FAKE_MARKET = {
    "Neighborhood Data Source": "Market Dashboard: test comp",
    "Future Occ/New/Canc": {
        "Category": {
            "test": {
                "Listings Used": 5,
                "X_values": ["2026-09-05", "2026-09-06"],
                "Y_values": [[70, 70], [0, 0], [0, 0], [40, 40], [40, 40], [0, 0]],
            }
        }
    },
    "Future Percentile Prices": {
        "Category": {
            "test": {
                "Listings Used": 5,
                "X_values": ["2026-09-05", "2026-09-06"],
                "Y_values": [[250, 250], [300, 300], [380, 380], [320, 320], [420, 420], [5, 5]],
            }
        }
    },
}

LISTINGS = [
    Listing(name="Test Property", pms="testpms", listing_id="id-1")
]

RESERVATION_ID_COL_0IDX = 29  # AD
STATUS_COL_0IDX = 28  # AC
COMP_CHECK_COL_0IDX = MANUAL_COLS_START - 1 + MANUAL_COLUMNS.index("Comp Check (Airbnb)")
FINAL_PL_CHECK_COL_0IDX = MANUAL_COLS_START - 1 + MANUAL_COLUMNS.index("Final PL Check")


def _reservation_row(res_id, checkin, checkout, confirmation_code, status="booked"):
    nights = (dt.date.fromisoformat(checkout) - dt.date.fromisoformat(checkin)).days
    return {
        "reservation_id": res_id,
        "listing_name": "Test Property",
        "check_in": checkin,
        "check_out": checkout,
        "booking_status": status,
        "rental_revenue": str(300 * nights),
        "no_of_days": nights,
        "booked_date": "2026-08-20",
        "booking_channel": "airbnb",
        "channelConfirmationCode": confirmation_code,
    }


def _make_client(reservations):
    client = MagicMock()
    client.get_reservations.return_value = reservations
    client.get_neighborhood_data.return_value = FAKE_MARKET
    client.get_overrides.return_value = []
    return client


def _run(tmp_path, today, reservations):
    with patch("booking_quality_log.main.cfg.load_listings", return_value=LISTINGS), \
         patch("booking_quality_log.main.cfg.get_api_key", return_value="fake"), \
         patch("booking_quality_log.main.PriceLabsClient", return_value=_make_client(reservations)):
        return bql_main.run(tmp_path, today=today)


def test_first_run_always_writes(tmp_path):
    res = [_reservation_row("r1", "2026-09-05", "2026-09-07", "AAAA111111")]
    wrote = _run(tmp_path, dt.date(2026, 9, 15), res)
    assert wrote is True
    assert len(list(tmp_path.glob("*.xlsx"))) == 1


def test_no_new_booking_skips_even_after_a_multi_day_gap(tmp_path):
    res = [_reservation_row("r1", "2026-09-05", "2026-09-07", "AAAA111111")]
    _run(tmp_path, dt.date(2026, 9, 15), res)

    wrote = _run(tmp_path, dt.date(2026, 9, 18), res)  # same booking, 3 days later
    assert wrote is False
    assert len(list(tmp_path.glob("*.xlsx"))) == 1


def test_new_booking_writes_and_carries_manual_notes_forward(tmp_path):
    res_day1 = [_reservation_row("r1", "2026-09-05", "2026-09-07", "AAAA111111")]
    _run(tmp_path, dt.date(2026, 9, 15), res_day1)

    [day1_file] = tmp_path.glob("*.xlsx")
    wb = openpyxl.load_workbook(day1_file)
    ws = wb["Booking Quality Log"]
    ws.cell(row=5, column=COMP_CHECK_COL_0IDX + 1, value="checked, price is fair")
    wb.save(day1_file)

    res_day3 = res_day1 + [_reservation_row("r2", "2026-09-20", "2026-09-22", "BBBB222222")]
    wrote = _run(tmp_path, dt.date(2026, 9, 20), res_day3)
    assert wrote is True

    files = sorted(tmp_path.glob("*.xlsx"))
    assert len(files) == 2

    wb2 = openpyxl.load_workbook(files[-1])
    ws2 = wb2["Booking Quality Log"]
    rows = list(ws2.iter_rows(min_row=5, values_only=True))
    ids = {r[RESERVATION_ID_COL_0IDX] for r in rows}
    assert ids == {"AAAA111111", "BBBB222222"}

    carried = next(r for r in rows if r[RESERVATION_ID_COL_0IDX] == "AAAA111111")
    assert carried[COMP_CHECK_COL_0IDX] == "checked, price is fair"


def test_cancellation_alone_does_not_trigger_a_new_file(tmp_path):
    res = [
        _reservation_row("r1", "2026-09-05", "2026-09-07", "AAAA111111"),
        _reservation_row("r2", "2026-09-10", "2026-09-11", "BBBB222222"),
    ]
    _run(tmp_path, dt.date(2026, 9, 15), res)

    # r2 comes back cancelled -- no brand-new *confirmed* booking appeared,
    # so per Thomas's choice this must still skip (even though r2 will now
    # show up as a Cancelled row once a file IS eventually written).
    res_after_cancellation = [
        _reservation_row("r1", "2026-09-05", "2026-09-07", "AAAA111111"),
        _reservation_row("r2", "2026-09-10", "2026-09-11", "BBBB222222", status="cancelled"),
    ]
    wrote = _run(tmp_path, dt.date(2026, 9, 16), res_after_cancellation)
    assert wrote is False
    assert len(list(tmp_path.glob("*.xlsx"))) == 1


def test_cancelled_booking_stays_visible_with_status_and_keeps_its_note(tmp_path):
    res_day1 = [
        _reservation_row("r1", "2026-09-05", "2026-09-07", "AAAA111111"),
        _reservation_row("r2", "2026-09-10", "2026-09-11", "BBBB222222"),
    ]
    _run(tmp_path, dt.date(2026, 9, 15), res_day1)

    [day1_file] = tmp_path.glob("*.xlsx")
    wb = openpyxl.load_workbook(day1_file)
    ws = wb["Booking Quality Log"]
    for row in ws.iter_rows(min_row=5):
        if row[RESERVATION_ID_COL_0IDX].value == "BBBB222222":
            row[FINAL_PL_CHECK_COL_0IDX].value = "priced well, sorry to lose it"
    wb.save(day1_file)

    # r2 is now cancelled, and r3 is a brand-new confirmed booking (needed
    # to actually trigger a write, since a cancellation alone never does).
    res_day2 = [
        _reservation_row("r1", "2026-09-05", "2026-09-07", "AAAA111111"),
        _reservation_row("r2", "2026-09-10", "2026-09-11", "BBBB222222", status="cancelled"),
        _reservation_row("r3", "2026-09-25", "2026-09-27", "CCCC333333"),
    ]
    wrote = _run(tmp_path, dt.date(2026, 9, 16), res_day2)
    assert wrote is True

    files = sorted(tmp_path.glob("*.xlsx"))
    wb2 = openpyxl.load_workbook(files[-1])
    ws2 = wb2["Booking Quality Log"]
    rows = list(ws2.iter_rows(min_row=5, values_only=True))

    cancelled_row = next(r for r in rows if r[RESERVATION_ID_COL_0IDX] == "BBBB222222")
    assert cancelled_row[STATUS_COL_0IDX] == "Cancelled"
    assert cancelled_row[FINAL_PL_CHECK_COL_0IDX] == "priced well, sorry to lose it"

    confirmed_row = next(r for r in rows if r[RESERVATION_ID_COL_0IDX] == "AAAA111111")
    assert confirmed_row[STATUS_COL_0IDX] == "Confirmed"


def test_rows_sorted_by_booked_date_newest_first(tmp_path):
    res = [
        _reservation_row("r1", "2026-09-05", "2026-09-07", "OLD111"),  # booked_date 2026-08-20 (fixed in helper)
    ]
    res[0]["booked_date"] = "2026-08-01"
    res.append(
        {**_reservation_row("r2", "2026-09-10", "2026-09-11", "NEW222"), "booked_date": "2026-09-14"}
    )
    wrote = _run(tmp_path, dt.date(2026, 9, 15), res)
    assert wrote is True

    [file] = tmp_path.glob("*.xlsx")
    wb = openpyxl.load_workbook(file)
    ws = wb["Booking Quality Log"]
    ids_in_order = [row[RESERVATION_ID_COL_0IDX] for row in ws.iter_rows(min_row=5, values_only=True)]
    assert ids_in_order == ["NEW222", "OLD111"]


def test_force_writes_even_with_no_new_booking(tmp_path):
    res = [_reservation_row("r1", "2026-09-05", "2026-09-07", "AAAA111111")]
    _run(tmp_path, dt.date(2026, 9, 15), res)

    with patch("booking_quality_log.main.cfg.load_listings", return_value=LISTINGS), \
         patch("booking_quality_log.main.cfg.get_api_key", return_value="fake"), \
         patch("booking_quality_log.main.PriceLabsClient", return_value=_make_client(res)):
        wrote = bql_main.run(tmp_path, today=dt.date(2026, 9, 18), force=True)
    assert wrote is True
    assert len(list(tmp_path.glob("*.xlsx"))) == 2


def test_override_reasons_pulled_into_the_sheet(tmp_path):
    res = [_reservation_row("r1", "2026-09-05", "2026-09-07", "AAAA111111")]
    client = _make_client(res)
    client.get_overrides.return_value = [
        {"date": "2026-09-05", "reason": "9/1 - Pacing behind by -10%"},
    ]
    with patch("booking_quality_log.main.cfg.load_listings", return_value=LISTINGS), \
         patch("booking_quality_log.main.cfg.get_api_key", return_value="fake"), \
         patch("booking_quality_log.main.PriceLabsClient", return_value=client):
        bql_main.run(tmp_path, today=dt.date(2026, 9, 15))

    [file] = tmp_path.glob("*.xlsx")
    wb = openpyxl.load_workbook(file)
    ws = wb["Booking Quality Log"]
    row = next(ws.iter_rows(min_row=5, values_only=True))
    override_notes_col_0idx = MANUAL_COLS_START - 2  # column right before the manual block
    assert row[override_notes_col_0idx] == "9/1 - Pacing behind by -10%"
